# -*- coding: utf-8 -*-
"""
Classification Gauche vs Droite avec CSP + LDA.
  • Métriques  : Accuracy + AUC-ROC
  • Stats      : Test de Wilcoxon signé (avant vs après ICA)
  • Export     : Excel  +  graphique
"""

import pickle
import os
import numpy as np
import matplotlib.pyplot as plt
import mne
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from sklearn.pipeline import make_pipeline
from sklearn.discriminant_analysis import LinearDiscriminantAnalysis as LDA
from sklearn.model_selection import StratifiedKFold, cross_val_score
from mne.decoding import CSP
from scipy.stats import wilcoxon

# ─────────────────────────────────────────────────────────────────────────────
# 1. CHEMINS
# ─────────────────────────────────────────────────────────────────────────────
base_path  = "data/epochs"
output_dir =  "results"

# ─────────────────────────────────────────────────────────────────────────────
# 2. CHARGEMENT
# ─────────────────────────────────────────────────────────────────────────────
print("Chargement des données Gauche vs Droite...")
with open(os.path.join(base_path, 'data_before.pkl'), 'rb') as f:
    data_before = pickle.load(f)
with open(os.path.join(base_path, 'data_after.pkl'), 'rb') as f:
    data_after = pickle.load(f)

# ─────────────────────────────────────────────────────────────────────────────
# 3. PIPELINE  CSP → LDA
# ─────────────────────────────────────────────────────────────────────────────
csp = CSP(n_components=4, reg='shrinkage', log=True, norm_trace=False)
lda = LDA()
clf = make_pipeline(csp, lda)

cv = StratifiedKFold(n_splits=5, shuffle=True, random_state=42)

# ─────────────────────────────────────────────────────────────────────────────
# 4. BOUCLE D'ÉVALUATION
# ─────────────────────────────────────────────────────────────────────────────
print("\n=== RÉSULTATS CSP + LDA — GAUCHE VS DROITE ===\n")

sujets     = list(data_before.keys())
acc_before = []
acc_after  = []
auc_before = []
auc_after  = []

for i in sujets:
    X_b, y_b = data_before[i]['X'], data_before[i]['y']
    X_a, y_a = data_after[i]['X'],  data_after[i]['y']

    sc_b_acc = cross_val_score(clf, X_b, y_b, cv=cv, scoring='accuracy', n_jobs=1)
    sc_a_acc = cross_val_score(clf, X_a, y_a, cv=cv, scoring='accuracy', n_jobs=1)
    sc_b_auc = cross_val_score(clf, X_b, y_b, cv=cv, scoring='roc_auc',  n_jobs=1)
    sc_a_auc = cross_val_score(clf, X_a, y_a, cv=cv, scoring='roc_auc',  n_jobs=1)

    mb_acc = np.mean(sc_b_acc) * 100;  ma_acc = np.mean(sc_a_acc) * 100
    mb_auc = np.mean(sc_b_auc);        ma_auc = np.mean(sc_a_auc)

    acc_before.append(mb_acc);  acc_after.append(ma_acc)
    auc_before.append(mb_auc);  auc_after.append(ma_auc)

    print(f"Sujet {i:02d} | Acc : {mb_acc:.1f}% → {ma_acc:.1f}%   "
          f"| AUC : {mb_auc:.3f} → {ma_auc:.3f}")

# ─────────────────────────────────────────────────────────────────────────────
# 5. TEST DE WILCOXON
# ─────────────────────────────────────────────────────────────────────────────
print("\n=== TEST STATISTIQUE (Wilcoxon signé, accuracy) ===")
stat, p = wilcoxon(acc_before, acc_after)
print(f"  W = {stat:.2f}  |  p = {p:.4f}")
if p < 0.05:
    print("  → Différence statistiquement SIGNIFICATIVE (p < 0.05)")
else:
    print("  → Différence NON significative (p ≥ 0.05)")

# ─────────────────────────────────────────────────────────────────────────────
# 6. EXPORT EXCEL
# ─────────────────────────────────────────────────────────────────────────────
print("\nSauvegarde du tableau Excel...")

# Styles
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
BEFORE_FILL = PatternFill("solid", fgColor="2E75B6")
AFTER_FILL  = PatternFill("solid", fgColor="70AD47")
DELTA_FILL  = PatternFill("solid", fgColor="D6E4F7")
MEAN_FILL   = PatternFill("solid", fgColor="FFD700")
WHITE_FONT  = Font(color="FFFFFF", bold=True)
BLACK_BOLD  = Font(bold=True)
CENTER      = Alignment(horizontal="center")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LDA_GaucheDroite"

# En-têtes
# Colonnes
headers = [
    "Sujet",
    "Accuracy AVANT ICA (%)",
    "Accuracy APRÈS ICA (%)",
    "Δ Accuracy",
    "AUC AVANT ICA",
    "AUC APRÈS ICA",
]
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill      = HEADER_FILL
    cell.font      = WHITE_FONT
    cell.alignment = CENTER

# Données par sujet
for row, i in enumerate(sujets, 2):
    delta = acc_after[row - 2] - acc_before[row - 2]

    ws.cell(row=row, column=1, value=f"Sujet {i:02d}")

    c_ab = ws.cell(row=row, column=2, value=round(acc_before[row - 2], 1))
    c_aa = ws.cell(row=row, column=3, value=round(acc_after[row - 2],  1))
    c_da = ws.cell(row=row, column=4, value=round(delta, 1))
    c_ub = ws.cell(row=row, column=5, value=round(auc_before[row - 2], 3))
    c_ua = ws.cell(row=row, column=6, value=round(auc_after[row - 2],  3))

    c_ab.fill = BEFORE_FILL; c_ab.font = Font(color="FFFFFF"); c_ab.alignment = CENTER
    c_aa.fill = AFTER_FILL;  c_aa.font = Font(color="FFFFFF"); c_aa.alignment = CENTER
    c_da.fill = DELTA_FILL;  c_da.alignment = CENTER
    c_ub.fill = BEFORE_FILL; c_ub.font = Font(color="FFFFFF"); c_ub.alignment = CENTER
    c_ua.fill = AFTER_FILL;  c_ua.font = Font(color="FFFFFF"); c_ua.alignment = CENTER

# Ligne de moyenne
mean_row = len(sujets) + 2
ws.cell(row=mean_row, column=1, value="MOYENNE").font = BLACK_BOLD
ws.cell(row=mean_row, column=1).fill = MEAN_FILL
ws.cell(row=mean_row, column=1).alignment = CENTER

mean_vals = [
    (2, round(np.mean(acc_before), 1)),
    (3, round(np.mean(acc_after),  1)),
    (4, round(np.mean(acc_after) - np.mean(acc_before), 1)),
    (5, round(np.mean(auc_before), 3)),
    (6, round(np.mean(auc_after),  3)),
]
for col, val in mean_vals:
    cell = ws.cell(row=mean_row, column=col, value=val)
    cell.fill = MEAN_FILL
    cell.font = BLACK_BOLD
    cell.alignment = CENTER

# Test de Wilcoxon
wil_row = mean_row + 2
ws.cell(row=wil_row, column=1, value="Test de Wilcoxon (avant vs après ICA)").font = BLACK_BOLD
ws.cell(row=wil_row, column=2, value=f"p = {p:.4f}").alignment = CENTER
sig = "Significatif (p < 0.05)" if p < 0.05 else "Non significatif"
ws.cell(row=wil_row, column=3, value=sig).font = BLACK_BOLD

# Largeur des colonnes
col_widths = [20, 24, 24, 14, 18, 18]
for col, width in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(col)].width = width

excel_path = os.path.join(output_dir, 'Résult_LDA_GaucheDroite.xlsx')
wb.save(excel_path)
print(f"  → {excel_path}")

# ─────────────────────────────────────────────────────────────────────────────
# 7. GRAPHIQUE COMPARATIF
# ─────────────────────────────────────────────────────────────────────────────
fig, ax = plt.subplots(figsize=(14, 5))
x     = np.arange(len(sujets))
width = 0.35

ax.bar(x - width/2, acc_before, width,
       label='Avant ICA', color='tomato',   alpha=0.85, edgecolor='white')
ax.bar(x + width/2, acc_after,  width,
       label='Après ICA',  color='steelblue', alpha=0.85, edgecolor='white')

ax.axhline(50, color='black', linestyle='--', linewidth=1.2, label='Chance (50 %)')
ax.set_xlabel('Sujet', fontsize=11)
ax.set_ylabel('Accuracy (%)', fontsize=11)
ax.set_title('CSP + LDA — Gauche vs Droite\nImpact de la correction ICA', fontsize=13)
ax.set_xticks(x)
ax.set_xticklabels([f"S{i:02d}" for i in sujets], rotation=45, ha='right')
ax.set_ylim(35, 100)
ax.legend(fontsize=10)

sig_label = f"Wilcoxon : p = {p:.4f}" + (" *" if p < 0.05 else "")
ax.text(0.98, 0.97, sig_label, transform=ax.transAxes,
        ha='right', va='top', fontsize=9,
        bbox=dict(boxstyle='round,pad=0.3', facecolor='lightyellow', edgecolor='gray'))

plt.tight_layout()
png_path = os.path.join(output_dir, 'Résult_LDA_GaucheDroite.png')
plt.savefig(png_path, dpi=150)
plt.show()
print(f"  → {png_path}")
